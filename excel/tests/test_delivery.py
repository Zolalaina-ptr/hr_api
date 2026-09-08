"""Contrôles de livraison portables. Ne compilent et n'exécutent PAS le VBA."""
import hashlib
import json
import re
import unittest
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from zipfile import ZipFile
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]

class DeliveryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(ROOT/'examples/GestionnaireRH-modele.xlsx')
        cls.forms = json.loads((ROOT/'tools/forms.json').read_text(encoding='utf-8'))
        cls.sources = {p.stem: p.read_text(encoding='utf-8') for p in (ROOT/'src').iterdir() if p.suffix in ('.bas','.vba')}
        cls.manifest = json.loads((ROOT/'tools/workbook-data.json').read_text(encoding='utf-8'))
        cls.tables = {t['name'] for sheet in cls.manifest['sheets'] for t in sheet['tables']}

    def test_required_sheets(self):
        self.assertTrue({'Accueil','Base_Salaries','Gestion_Conges','Parametres','Documents','Dashboard_Stat','Admin_Logs'} <= set(self.wb.sheetnames))

    def test_only_home_visible(self):
        self.assertEqual([s.title for s in self.wb if s.sheet_state=='visible'], ['Accueil'])

    def test_five_unique_fictitious_employees(self):
        ws = self.wb['Base_Salaries']
        self.assertEqual(ws.max_row, 6)
        ids = [ws.cell(r,1).value for r in range(2,7)]
        self.assertEqual(ids, [f'SAL{i:06}' for i in range(1,6)])
        for r in range(2,7):
            self.assertTrue(ws.cell(r,8).value.endswith('@example.invalid'))
            self.assertIsNone(ws.cell(r,6).value)
            self.assertIsNone(ws.cell(r,15).value)
            if ws.cell(r,17).value: self.assertIn(ws.cell(r,17).value, ids)

    def test_employee_column_contract(self):
        headers=[c.value for c in self.wb['Base_Salaries'][1]]
        self.assertEqual(len(headers),22)
        for index,name in {1:'Matricule',5:'Naissance',10:'Embauche',17:'Manager',18:'Statut',19:'Sortie',21:'SoldeInitial',22:'Sexe'}.items():
            self.assertEqual(headers[index-1],name)

    def test_leave_column_contract(self):
        self.assertEqual([c.value for c in self.wb['Gestion_Conges'][1]], ['Identifiant','Matricule','Debut','Fin','Type','Jours','Statut','Demandeur','Creation','Validateur','Decision','Motif'])

    def test_all_literal_table_references_exist(self):
        for source in self.sources.values():
            for name in re.findall(r'(?:Table|FindRow)\("([^"]+)"',source):
                self.assertIn(name,self.tables)

    def test_all_literal_sheet_references_exist(self):
        for source in self.sources.values():
            for name in re.findall(r'Worksheets\("([^"]+)"\)',source):
                self.assertIn(name,self.wb.sheetnames)

    def test_option_explicit_and_documented_procedures(self):
        for name,source in self.sources.items():
            with self.subTest(module=name):
                self.assertIn('Option Explicit',source)
                lines=source.splitlines()
                for i,line in enumerate(lines):
                    self.assertLessEqual(len(line),1023)
                    if re.match(r'(Public |Private )?(Sub|Function) ',line):
                        self.assertIn('Auteur : Arena | Date : 2026-09-08 | Description :', lines[i-1])

    def test_form_names_and_events_match(self):
        for form in self.forms:
            source=self.sources[form['name']]
            controls={c['name'] for c in form['controls']}
            self.assertEqual(len(controls),len(form['controls']))
            for control in form['controls']:
                if control['type']=='CommandButton':
                    self.assertIn('Private Sub '+control['name']+'_Click()',source)
            for name in re.findall(r'Me\.([A-Za-z0-9_]+)',source):
                self.assertTrue(name in controls or name=='Controls',name)

    def test_employee_form_has_all_22_fields(self):
        form=next(f for f in self.forms if f['name']=='frmEmployee')
        names={c['name'] for c in form['controls']}
        self.assertTrue({f'f{i}' for i in range(1,23)} <= names)
        self.assertTrue(next(c for c in form['controls'] if c['name']=='f1')['locked'])

    def test_password_controls_masked(self):
        for name in ['frmLogin','frmAccount']:
            form=next(f for f in self.forms if f['name']==name)
            self.assertEqual(next(c for c in form['controls'] if c['name']=='password')['passwordChar'],'*')

    def test_no_default_password(self):
        ws=self.wb['Utilisateurs']
        self.assertIsNone(ws.cell(2,4).value)
        self.assertEqual(ws.cell(2,5).value,'A_INITIALISER')
        self.assertIn('Read-Host',(ROOT/'tools/Build-Workbook.ps1').read_text(encoding='utf-8-sig'))

    def test_crypto_cross_language_fixture(self):
        expected=hashlib.sha256('abc'.encode('utf-16le')).hexdigest()
        self.assertIn(expected,self.sources['modTests'])
        self.assertIn('[Text.Encoding]::Unicode',(ROOT/'tools/Build-Workbook.ps1').read_text(encoding='utf-8-sig'))

    def test_menu_targets_exist(self):
        ps=(ROOT/'tools/Build-Workbook.ps1').read_text(encoding='utf-8-sig')
        menu=ps.split('$buttons = @(',1)[1].split('$index = 0',1)[0]
        public_procedures=set(re.findall(r'Public Sub (\w+)\(', '\n'.join(self.sources.values())))
        targets=re.findall(r"@\('[^']+','([^']+)'\)",menu)
        self.assertEqual(len(targets),13)
        self.assertTrue(set(targets)<=public_procedures)

    def test_no_broken_template_tags(self):
        sheet=next(s for s in self.manifest['sheets'] if s['name']=='Parametres')
        table=next(t for t in sheet['tables'] if t['name']=='Modeles')
        allowed={'Entreprise','Nom','Prenom','Poste','Contrat','Salaire','Embauche','Sortie','Date'}
        for row in self.wb['Parametres'][table['range']][1:]:
            self.assertTrue(set(re.findall(r'\[([^]]+)\]',row[1].value))<=allowed)
            self.assertIn('BROUILLON',row[1].value)

    def test_manual_eight_source_sections_and_pdf_pages(self):
        self.assertEqual(len((ROOT/'docs/manuel-utilisateur.md').read_text().split('\n---\n')),8)
        pdf=(ROOT/'docs/Manuel-utilisateur.pdf').read_bytes()
        self.assertTrue(pdf.startswith(b'%PDF-'))
        self.assertEqual(len(re.findall(rb'/Type\s*/Page\b',pdf)),8)

    def test_saved_views_are_cleared(self):
        self.assertIn('LockWorkbook',self.sources['ThisWorkbook'])
        self.assertIn('chart.Delete',self.sources['modCore'])
        self.assertIn('xlSheetVeryHidden',self.sources['modCore'])
        self.assertIn('Cancel = True',self.sources['ThisWorkbook'])

    def test_export_has_scope_check(self):
        code=self.sources['modReports'].split('Public Sub ExportEmployees()',1)[1]
        self.assertIn('CanRead',code)
        self.assertNotIn('data(i, 6)',code)  # NSS
        self.assertNotIn('data(i, 14)',code) # Salaire
        self.assertNotIn('data(i, 15)',code) # RIB

    def test_handlers_for_each_form_event(self):
        for form in self.forms:
            source=self.sources[form['name']]
            self.assertEqual(source.count('Private Sub '),source.count('On Error GoTo Failed'))

    def test_windows_api_uses_pointer_safe_declarations(self):
        declarations=[l for l in self.sources['modSecurity'].splitlines() if 'Declare ' in l]
        self.assertEqual(len(declarations),6)
        self.assertTrue(all('PtrSafe' in l for l in declarations))
        self.assertTrue(all('LongPtr' in l for l in declarations))


    def test_native_build_does_not_open_or_repair_xlsx(self):
        builder=(ROOT/'tools/Build-Workbook.ps1').read_text(encoding='utf-8-sig')
        self.assertNotIn('.Workbooks.Open(', builder)
        self.assertNotIn('CorruptLoad', builder)
        self.assertIn('.Workbooks.Add(-4167)', builder)
        self.assertIn('Initialize-RHWorkbook -Workbook $book -Manifest $manifest', builder)

    def test_preview_contains_no_ooxml_tables(self):
        with ZipFile(ROOT/'examples/GestionnaireRH-modele.xlsx') as archive:
            self.assertFalse(any(n.startswith('xl/tables/') for n in archive.namelist()))
            for name in archive.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    self.assertNotIn(b'<tableParts', archive.read(name))

    def test_manifest_matches_all_preview_values(self):
        self.assertEqual(self.manifest['schemaVersion'],1)
        self.assertEqual([s['name'] for s in self.manifest['sheets']], self.wb.sheetnames)
        for sheet in self.manifest['sheets']:
            actual={c.coordinate:c.value for row in self.wb[sheet['name']] for c in row if c.value is not None}
            expected={}
            for cell in sheet['cells']:
                self.assertNotIn(cell['address'],expected)
                self.assertIn(cell['type'],('date','number','text'))
                value=cell['value']
                if cell['type']=='date': value=datetime.strptime(value,'%Y-%m-%d')
                expected[cell['address']]=value
            self.assertEqual(actual,expected,sheet['name'])

    def test_native_tables_have_seed_row_and_exact_headers(self):
        names=[]
        for sheet in self.manifest['sheets']:
            occupied=set()
            for table in sheet['tables']:
                names.append(table['name'])
                left,top,right,bottom=range_boundaries(table['nativeRange'])
                self.assertGreater(bottom,top)
                self.assertEqual(bottom-top,max(1,table['dataRows']))
                self.assertEqual(table['headers'],[self.wb[sheet['name']].cell(top,c).value for c in range(left,right+1)])
                cells={(r,c) for r in range(top,bottom+1) for c in range(left,right+1)}
                self.assertFalse(occupied & cells)
                occupied |= cells
        self.assertEqual(len(names),11)
        self.assertEqual(len(names),len(set(names)))

    def test_native_builder_removes_empty_seed_rows_and_checks_schema(self):
        code=(ROOT/'tools/Initialize-Workbook.ps1').read_text(encoding='utf-8-sig')
        self.assertIn('$table.ListRows.Item(1).Delete()',code)
        self.assertIn('$table.ListRows.Count -ne [int]$definition.dataRows',code)
        self.assertIn('$table.ListColumns.Count -ne $definition.headers.Count',code)
        for sheet in self.manifest['sheets']:
            for table in sheet['tables']:
                if table['name'] in ('Conges','Logs','Competences','Entretiens'):
                    self.assertEqual(table['dataRows'],0)

    def test_builder_errors_name_stage_and_cleanup_is_guarded(self):
        code=(ROOT/'tools/Build-Workbook.ps1').read_text(encoding='utf-8-sig')
        self.assertIn('$stage = "Création du classeur vierge"',code)
        self.assertIn('InvocationInfo.ScriptLineNumber',code)
        self.assertIn('try { $book.Close($false) } catch',code)
        self.assertIn('try { $excel.Quit() } catch',code)

if __name__=='__main__':
    unittest.main(verbosity=2)
