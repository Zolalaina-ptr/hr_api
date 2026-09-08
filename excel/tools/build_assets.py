"""Auteur : Arena | 2026-09-08. Modèle sans macro, paramètres et manuel reproductibles.
Exécution : python excel/tools/build_assets.py (openpyxl, reportlab requis).
Aucun vrai dossier personnel n'est inclus. Le .xlsm est assemblé sous Windows.
"""
from datetime import datetime
from pathlib import Path
import csv
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
NAVY = '142D4E'
TABLES = {
    'Base_Salaries': ('Salaries', ['Matricule','Civilite','Nom','Prenom','Naissance','NSS','Telephone','Email','Adresse','Embauche','Contrat','Service','Poste','Salaire','RIB','Photo','Manager','Statut','Sortie','Motif','SoldeInitial','Sexe']),
    'Gestion_Conges': ('Conges', ['Identifiant','Matricule','Debut','Fin','Type','Jours','Statut','Demandeur','Creation','Validateur','Decision','Motif']),
    'Utilisateurs': ('Utilisateurs', ['Utilisateur','Role','Matricule','Sel','Empreinte','Actif']),
    'Admin_Logs': ('Logs', ['Horodatage','Utilisateur','Action','Cible','SessionWindows']),
    'Competences': ('Competences', ['Identifiant','Matricule','Competence','Niveau','Date']),
    'Entretiens': ('Entretiens', ['Identifiant','Matricule','Date','Remarques','Auteur']),
}
CONFIG = [
    ['Entreprise', 'Entreprise Démonstration'], ['AcquisitionMensuelle', 2.5],
    ['ModeJours', 'Ouvres'], ['DebutExercice', '01/01/2026'],
    ['CouleurMarine', 5123348], ['DossierExports', ''],
    ['Protection', 'INSTALLATION'], ['CompteurSalaries', 5],
    ['CompteurConges', 0], ['CompteurEntretiens', 0],
]
REFERENCES = {
    'Service': ['Direction', 'Finance', 'Ressources humaines', 'Informatique'],
    'Poste': ['Directeur', 'Comptable', 'Responsable RH', 'Développeur'],
    'Contrat': ['CDI','CDD','Intérim'],
    'Motif': ['Démission','Licenciement','Retraite','Fin de CDD'],
    'Conge': ['Payé','Sans solde','RTT','Maladie'],
}
MODELS = [
    ['Contrat', 'BROUILLON À VALIDER JURIDIQUEMENT\nCONTRAT DE TRAVAIL\n\nEntre [Entreprise] et [Prenom] [Nom].\nType : [Contrat]. Poste : [Poste].\nDate d’embauche : [Embauche]. Salaire brut : [Salaire] (devise et périodicité à préciser).\n\nCe modèle ne constitue pas un contrat complet : compléter convention collective, durée du travail, lieu, période d’essai, clauses et signatures.\nÉtabli le [Date].'],
    ['Avenant', 'BROUILLON À VALIDER JURIDIQUEMENT\nAVENANT\n\n[Entreprise] / [Prenom] [Nom]\nPoste actuel : [Poste]. Salaire brut actuel : [Salaire].\nCompléter la modification, sa date d’effet et les clauses inchangées avant signature.\nÉtabli le [Date].'],
    ['Certificat', 'BROUILLON À VALIDER\nCERTIFICAT DE TRAVAIL\n\n[Entreprise] certifie avoir employé [Prenom] [Nom] du [Embauche] au [Sortie] au poste de [Poste].\nCompléter les mentions légales applicables et la signature du représentant.\nÉtabli le [Date].'],
    ['Attestation', 'BROUILLON À VALIDER\nATTESTATION\n\n[Entreprise] atteste que [Prenom] [Nom] a été embauché(e) le [Embauche] au poste de [Poste].\nCompléter l’objet et vérifier la situation actuelle avant signature.\nÉtabli le [Date].'],
]

def add_table(ws, name, headers, rows, col=1):
    for j,h in enumerate(headers,col): ws.cell(1,j,h)
    for i,row in enumerate(rows,2):
        for j,value in enumerate(row,col):
            if value != '': ws.cell(i,j,value)
    from openpyxl.utils import get_column_letter
    tab = Table(displayName=name, ref=f'{get_column_letter(col)}1:{get_column_letter(col+len(headers)-1)}{max(1,len(rows)+1)}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tab)

def employees():
    records = [
        ('Mme','Martin','Alice','1985-09-12','2020-01-01','Direction','Directeur','','F'),
        ('M.','Bernard','Lucas','1990-03-04','2022-01-01','Finance','Comptable','SAL000001','H'),
        ('Mme','Petit','Emma','1993-09-24','2023-01-01','Ressources humaines','Responsable RH','SAL000001','F'),
        ('M.','Robert','Hugo','1988-06-15','2021-01-01','Informatique','Développeur','SAL000001','H'),
        ('Mme','Durand','Chloé','1996-11-07','2024-01-01','Finance','Comptable','SAL000001','F'),
    ]
    for i,(title,last,first,birth,hire,service,post,manager,sex) in enumerate(records,1):
        yield [f'SAL{i:06}',title,last,first,datetime.fromisoformat(birth),'','',f'demo{i}@example.invalid','Adresse fictive',datetime.fromisoformat(hire),'CDI',service,post,2500,'','',manager,'Actif','','',5,sex]

def workbook():
    wb = Workbook(); home = wb.active; home.title='Accueil'
    home['B2']='RH / ESPACE DE GESTION'; home['D4']='Connectez-vous pour consulter votre espace RH.'
    home['D2']='Gestionnaire RH • Excel / VBA • Démonstration'
    home['D3']='Données fictives — un seul utilisateur à la fois — sauvegardez après vos opérations.'
    for sheet,(name,headers) in TABLES.items():
        ws = wb.create_sheet(sheet)
        rows = list(employees()) if name=='Salaries' else []
        if name=='Utilisateurs': rows=[['admin','Admin','','','A_INITIALISER','Oui']]
        add_table(ws,name,headers,rows)
    ws=wb.create_sheet('Parametres')
    add_table(ws,'Configuration',['Cle','Valeur'],CONFIG)
    add_table(ws,'References',['Categorie','Valeur'],[(k,v) for k,values in REFERENCES.items() for v in values],4)
    # Exemple France métropolitaine 2026, à adapter au site et à la convention.
    holidays=['2026-01-01','2026-04-06','2026-05-01','2026-05-08','2026-05-14','2026-05-25','2026-07-14','2026-08-15','2026-11-01','2026-11-11','2026-12-25']
    add_table(ws,'Feries',['Date','Libelle'],[(datetime.fromisoformat(d),'Exemple France 2026') for d in holidays],7)
    add_table(ws,'Grille',['Poste','Competence','NiveauCible'],[['Comptable','Excel','Expert'],['Comptable','Sage','Intermédiaire'],['Comptable','Compta Générale','Expert'],['Développeur','VBA','Expert'],['Responsable RH','Excel','Intermédiaire'],['Directeur','Management','Expert']],10)
    add_table(ws,'Modeles',['Modele','Texte'],MODELS,14)
    wb.create_sheet('Documents'); wb.create_sheet('Dashboard_Stat')
    for ws in wb:
        ws.sheet_view.showGridLines=False; ws.freeze_panes='A2'
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(name='Calibri',bold=True,color='FFFFFF')
        for col in range(1,max(ws.max_column,14)+1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col)].width=21
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font=Font(name='Calibri',size=11,color=NAVY)
                if isinstance(cell.value,datetime):cell.number_format='dd/mm/yyyy'
        if ws!=home:ws.sheet_state='veryHidden'
    home.column_dimensions['A'].width=3;home.column_dimensions['B'].width=27;home.column_dimensions['C'].width=3
    home['B2'].font=Font(name='Calibri',size=13,bold=True,color='FFFFFF')
    home['B2'].fill=PatternFill('solid',fgColor=NAVY)
    home['D2'].font=Font(name='Calibri',size=20,bold=True,color=NAVY)
    home.sheet_properties.pageSetUpPr.fitToPage=True
    # Le constructeur Windows utilise un manifeste de données et laisse Excel
    # créer ses propres ListObjects. Aucun Workbooks.Open sur un XLSX externe.
    manifest = {'schemaVersion': 1, 'sheets': []}
    for sheet in wb:
        spec = {'name': sheet.title, 'cells': [], 'tables': [], 'widths': []}
        for row in sheet:
            for cell in row:
                if cell.value is None:
                    continue
                value = cell.value
                kind = 'text'
                if isinstance(value, datetime):
                    value = value.strftime('%Y-%m-%d'); kind = 'date'
                elif isinstance(value, (int, float)):
                    kind = 'number'
                spec['cells'].append({'address': cell.coordinate, 'type': kind, 'value': value})
        for name in list(sheet.tables):
            table = sheet.tables[name]
            from openpyxl.utils.cell import range_boundaries
            left, top, right, bottom = range_boundaries(table.ref)
            spec['tables'].append({'name': name, 'range': table.ref,
                'nativeRange': f'{get_column_letter(left)}{top}:{get_column_letter(right)}{max(top+1,bottom)}',
                'dataRows': bottom-top,
                'headers': [sheet.cell(top,c).value for c in range(left,right+1)]})
            # Le XLSX devient un aperçu simple, sans parties table ni autoFilter.
            # Les véritables tables seront sérialisées nativement par Excel Windows.
            del sheet.tables[name]
        for column, dim in sheet.column_dimensions.items():
            spec['widths'].append({'column': column, 'width': dim.width})
        manifest['sheets'].append(spec)
    (ROOT/'tools/workbook-data.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    dest=ROOT/'examples/GestionnaireRH-modele.xlsx';wb.save(dest)
    with (ROOT/'examples/Parametres.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f,delimiter=';');writer.writerow(['Cle','Valeur']);writer.writerows([r for r in CONFIG if not r[0].startswith('Compteur') and r[0]!='Protection'])
    params=Workbook();params.remove(params.active)
    for name,headers,rows in [('Configuration',['Cle','Valeur'],[r for r in CONFIG if not r[0].startswith('Compteur') and r[0]!='Protection']),('References',['Categorie','Valeur'],[(k,v) for k,values in REFERENCES.items() for v in values]),('Feries',['Date','Libelle'],[(datetime.fromisoformat(d),'Exemple France 2026') for d in holidays]),('Grille',['Poste','Competence','NiveauCible'],[['Comptable','Excel','Expert']]),('Modeles',['Modele','Texte'],MODELS)]:
        sheet=params.create_sheet(name);sheet.append(headers)
        for row in rows:sheet.append(row)
        sheet.column_dimensions['A'].width=30;sheet.column_dimensions['B'].width=90
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,datetime):cell.number_format='dd/mm/yyyy'
    params.save(ROOT/'examples/Parametres-entreprise.xlsx')

def manual():
    sections=(ROOT/'docs/manuel-utilisateur.md').read_text(encoding='utf-8').split('\n---\n')
    styles=getSampleStyleSheet();styles['Heading1'].textColor=colors.HexColor('#142D4E')
    styles['BodyText'].fontSize=10;styles['BodyText'].leading=15
    story=[]
    for i,section in enumerate(sections):
        if i:story.append(PageBreak())
        for line in section.strip().splitlines():
            if not line.strip():continue
            heading=line.startswith('#')
            story.append(Paragraph(escape(line.lstrip('# ').strip()),styles['Heading1'] if heading else styles['BodyText']))
            story.append(Spacer(1,7 if heading else 3))
    def footer(canvas,doc):
        canvas.setFont('Helvetica',9);canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawString(42,28,'GESTIONNAIRE RH / Manuel v1.0 / Données fictives')
        canvas.drawRightString(550,28,str(doc.page))
    SimpleDocTemplate(str(ROOT/'docs/Manuel-utilisateur.pdf'),rightMargin=42,leftMargin=42,topMargin=40,bottomMargin=48).build(story,onFirstPage=footer,onLaterPages=footer)

if __name__=='__main__':
    workbook();manual()
